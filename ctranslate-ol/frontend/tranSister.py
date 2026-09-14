# -*- coding: utf-8 -*-
import os, time, sys, json, requests

def eprint(*args, **kwargs):
	print(*args, file=sys.stderr, **kwargs)

service_url = os.getenv("BACKEND_URL") #'http://188.34.158.8:25000'

for _ in range(30):
	try:
		info = requests.get(service_url + '/info', timeout=2)
		break
	except Exception:
		time.sleep(1)

models = json.loads(info.text).get('models')
directions = set(sum((model.get('directions') for model in models), []))
srcs, tgts = map(set, zip(*(dir.split('_') for dir in directions)))
notto = set(f'{src}_{tgt}' for src in srcs for tgt in tgts if src != tgt) - directions

from flask import Flask, Response, request, render_template, send_file, redirect
# from flask_cors import CORS
app = Flask(__name__)
# CORS(app)
app.config["DEBUG"] = False

@app.route('/translate', methods=['POST'])
def translate():
	# CORS erlauben
	headers = {"Access-Control-Allow-Origin": "*"}

	body = request.get_data()
	r = requests.post(service_url + '/translate', data=body, headers={"Content-Type": "application/json" })
	return Response(
        r.content,
        status=r.status_code,
        mimetype="application/json",
        headers=headers
    )

@app.route('/spell', methods=['POST'])
def spell():
	reqdata = request.get_json()
	from hunspell import Hunspell # pip install cyhunspell
	speller = Hunspell({'hsb':'hsb_DE', 'de':'de_DE', 'dsb':'dsb_DE', 'cs': 'cs_CZ'}[reqdata.get('lang')], hunspell_data_dir='dictionaries')
	ignored = str.maketrans('', '', '.›‹')

	def collect_errors():
		for word in reqdata.get('words'):
			if not speller.spell(word.translate(ignored)):
				yield (word, '\n'.join(speller.suggest(word)))

	return dict(collect_errors())

@app.route('/download', methods=['POST'])
def download_dokument():
	from docx import Document # pip install python-docx
	from docx.shared import Inches, Cm
	from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX # type: ignore
	from docx.enum.section import WD_ORIENT # type: ignore
	from docx.oxml import OxmlElement
	from docx.oxml.ns import qn

	import re

	def divide_str(text, keywords):
		if not keywords: return [text]
		return list(re.split('(\\b' + '\\b|\\b'.join(keywords) + '\\b)', text))

	def set_celltext(cell, text, marked, lang):
		cell_para = cell.paragraphs[0]
		for section in divide_str(text, marked):
			cell_run = cell_para.add_run()
			cell_rpr = cell_run.element.get_or_add_rPr()
			cell_lang = OxmlElement('w:lang')
			cell_lang.set(qn('w:val'), {'hsb':'hsb-DE', 'de':'de-DE', 'dsb':'dsb-DE', 'cs':'cs-CZ'}[lang])
			cell_rpr.append(cell_lang)
			cell_run.add_text(section)
			if section in marked:
				cell_run.font.highlight_color = WD_COLOR_INDEX.YELLOW # type: ignore

	def set_repeat_table_header(row):
		trPr = row._tr.get_or_add_trPr()
		tblHeader = OxmlElement('w:tblHeader')
		tblHeader.set(qn('w:val'), "true")
		trPr.append(tblHeader)
		for col in row.cells: # Schattierung hinzufügen
			tcPr = col._tc.get_or_add_tcPr()
			clShading = OxmlElement('w:shd')
			clShading.set(qn('w:fill'), "D3D3D3")
			tcPr.append(clShading)

	class column:
		def __init__(self, caption, content):
			self.caption = caption
			self.content = content

	reqdata = request.get_json()
	option, src, tgt = reqdata.get('option'), reqdata.get('from'), reqdata.get('to')
	input, output, unks, used_model = reqdata.get('box_data').values()

	columns = [
		column(src, [sentence for item in input  for sentence in item if len(sentence)]),
		column(tgt, [sentence for item in output for sentence in item if len(sentence)])
	]
	columns[1].caption += '\n' + used_model
	compare_models = set()

	A4height, A4width = Cm(29.7), Cm(21.0)
	marginwidth = Cm(2.0)
	numwidth = Cm(1.0)

	if option == 'multicol':
		compare_models = set(model.get('name') for model in models if f'{src}_{tgt}' in model.get('directions')) - { used_model }
		for model in compare_models:
			payload = { 'source_language': src, 'target_language': tgt, 'model': model, 'text': '\n'.join([' '.join(item) for item in input]) }
			response = requests.post(service_url + '/translate', json = payload)
			columns.append(column(f'{tgt}\n{model}', [sentence for item in response.json().get('marked_translation') for sentence in item if len(sentence)]))

	orientation, height, width = (WD_ORIENT.LANDSCAPE, A4width, A4height) if len(columns) > 2 else (WD_ORIENT.PORTRAIT, A4height, A4width)
	columnwidth = round((width - 2*marginwidth - numwidth) / float(len(columns)))

	# Dokument erstellen
	doc = Document()
	for section in doc.sections:
		section.orientation = orientation
		section.page_height = height
		section.page_width  = width
		section.top_margin, section.bottom_margin, section.left_margin, section.right_margin = (marginwidth,)*4

	# Tabelle erstellen
	table = doc.add_table(rows = 1, cols = len(columns) + 1)
	table.style = 'Table Grid'
	proof_err = OxmlElement('w:proofErr')
	proof_err.set(qn('w:type'), 'spellStart')
	table._element.append(proof_err)

	# Überschriftenzeile hinzufügen
	set_repeat_table_header(table.rows[0])
	hdr_cells = table.rows[0].cells

	for i, caption in enumerate(['#'] + [column.caption for column in columns]):
		table.columns[i].width = columnwidth if i else numwidth
		hdr_cells[i].width = table.columns[i].width
		hdr_cells[i].text = caption
		hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

	# Inhalt hinzufügen
	for line in zip(*[column.content for column in columns]):
		if option == 'unks_only' and len(divide_str(line[0] + line[1], unks)) == 1:
			continue
		tr = table.add_row()
		trPr = tr._tr.get_or_add_trPr()
		trPr.append(OxmlElement('w:cantSplit'))
		tr.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
		tr.cells[0].paragraphs[0].style = 'List Number'

		for col in range(len(columns)):
			set_celltext(tr.cells[col+1], line[col], unks if col<2 else [], src if col == 0 else tgt)

	proof_err = OxmlElement('w:proofErr')
	proof_err.set(qn('w:type'), 'spellEnd')
	table._element.append(proof_err)

	# Dokument speichern
	from io import BytesIO
	buffer = BytesIO()
	doc.save(buffer)
	buffer.seek(0)

	return send_file(buffer, as_attachment=True, download_name='Download.docx') # mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'

jinjas = {
	'lang':       ('de',                    'hsb',                          'dsb'),
	# =====================================================================================================
	'inlang':     ('Eingabesprache',        'rěč zapodaća',                 'rěc zapódaśa'),
	'outlang':    ('Ausgabesprache',        'rěč wudaća',                   'rěc wudaśa'),
	'translate':  ('Übersetzen',            'přełožić',                     'pśełožyś'),
	'showbreaks': ('Absatzmarken anzeigen', 'wotbytkowe znamjenja pokazać', 'wótstawkowe znamjenja pokazaś'),
	'spellcheck': ('Rechtschreibprüfung',   'prawopisna kontrola',          'pśespytowanje pšawopisa'),
	'unknown':    ('unbekannt',             'njeznate',                     'njeznate'),
	'clearx':     ('Eingabetext löschen',   'zapodaty tekst wotstronić',    'zapódany text wótpóraś'),
	'neutral':    ('Sprache erkennen',      'rěč spóznać',                  'rěč spóznaś'),
	'de':         ('Deutsch',               'němsce',                       'nimski'),
	'hsb':        ('Obersorbisch',          'hornjoserbsce',                'górnoserbšćinu'),
	'dsb':        ('Niedersorbisch',        'delnjoserbsce',                'dolnoserbšćinu'),
	'cs':         ('Tschechisch',           'čěsce',                        'čěsćinu'),

	'download':   ('.docx herunterladen',   '.docx sćahnyć',                '.docx wóśěgnuś'),
	'ttable':     ('Übersetzungstabelle',   'přełožowanska_tabulka',        'pśełožowańska_tabulka'),
	'twocol':     ('zweispaltig',           'dwušpaltowy',                  'dwušpaltowy (?)'),
	'multicol':   ('mehrspaltig',           'wjacešpaltowy',                'wjacešpaltowy (?)'),
	'unks_only':  ('nur unbekannte',        'jenož njeznate',               'jenož njeznaty (?)'),

	'footer':     ('Fußzeile',              'nóžka',                        'stojnica')
}

@app.route("/")
def root():
    return redirect("/hsb")

@app.route('/<par>') # type: ignore

# @cross_origin()
def deliver(par):
	param = str(par)
	if param == 'favicon.ico':      return send_file('templates/favicon.ico', mimetype='image/x-icon')
	if param == 'App.css':          return send_file('templates/App.css', mimetype='text/css')
	if param == 'guessLanguage.js': return send_file('templates/guessLanguage.js', mimetype='text/javascript')
	if param == 'App.js':           return render_template('App.js', info=info.text.strip())

	if param in jinjas['lang']:
		col = jinjas['lang'].index(param)
		return render_template('App.htm', **{k:v[col] for k, v in jinjas.items()})

if __name__ == '__main__':
	# s.a. https://stackoverflow.com/questions/51025893/flask-at-first-run-do-not-use-the-development-server-in-a-production-environmen
	# app.run('0.0.0.0', 5000, ssl_context='adhoc')
	app.run('0.0.0.0', 3000)
